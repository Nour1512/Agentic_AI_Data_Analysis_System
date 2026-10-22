import logging
from pathlib import Path
from agents_schema.report import StructuredReport

logger = logging.getLogger(__name__)

class ReportRenderer:
    """Renders a StructuredReport into HTML and DOCX formats."""

    @staticmethod
    def render_html(report: StructuredReport) -> str:
        """Render the structured report to semantic HTML for the preview viewer."""
        html = []
        html.append(f'<div class="report-document">')
        
        # Header
        html.append(f'<div class="report-header">')
        html.append(f'<h1 class="report-title">{report.title}</h1>')
        if report.subtitle:
            html.append(f'<h2 class="report-subtitle">{report.subtitle}</h2>')
        
        # Metadata
        html.append('<div class="report-metadata">')
        for k, v in report.metadata.items():
            html.append(f'<span><strong>{k.replace("_", " ").title()}:</strong> {v}</span>')
        html.append('</div>')
        html.append('</div>')  # end report-header
        
        # Executive Summary
        html.append('<div class="report-section">')
        html.append('<h3>Executive Summary</h3>')
        html.append(f'<p>{report.executive_summary}</p>')
        html.append('</div>')

        # Objective
        html.append('<div class="report-section">')
        html.append('<h3>Analysis Objective</h3>')
        html.append(f'<p>{report.objective}</p>')
        html.append('</div>')

        # Data Overview
        if report.data_overview:
            html.append('<div class="report-section">')
            html.append('<h3>Data Overview</h3>')
            html.append(f'<p>{report.data_overview}</p>')
            html.append('</div>')

        # KPIs
        if report.kpis:
            html.append('<div class="report-section">')
            html.append('<h3>Key Metrics</h3>')
            html.append('<div class="report-metrics-grid">')
            for kpi in report.kpis:
                html.append('<div class="report-metric-card">')
                html.append(f'<div class="metric-label">{kpi.label}</div>')
                html.append(f'<div class="metric-value">{kpi.value}</div>')
                if kpi.description:
                    html.append(f'<div class="metric-desc">{kpi.description}</div>')
                html.append('</div>')
            html.append('</div>')
            html.append('</div>')

        # Key Findings
        if report.key_findings:
            html.append('<div class="report-section">')
            html.append('<h3>Key Findings</h3>')
            html.append('<ul>')
            for finding in report.key_findings:
                html.append(f'<li>{finding}</li>')
            html.append('</ul>')
            html.append('</div>')

        # Visual Analysis (Charts)
        if report.visualizations:
            html.append('<div class="report-section">')
            html.append('<h3>Visual Analysis</h3>')
            for viz in report.visualizations:
                html.append('<div class="report-chart-container">')
                html.append(f'<h4>{viz.title}</h4>')
                html.append(f'<img src="/{viz.image_path}" alt="{viz.title}" class="report-image">')
                if viz.insight:
                    html.append(f'<p class="chart-caption">{viz.insight}</p>')
                html.append('</div>')
            html.append('</div>')

        # Insights
        if report.insights:
            html.append('<div class="report-section">')
            html.append('<h3>AI Insights</h3>')
            html.append('<div class="report-insights-grid">')
            for insight in report.insights:
                html.append('<div class="report-insight-card">')
                html.append(f'<div class="insight-title">{insight.title}</div>')
                if insight.value:
                    html.append(f'<div class="insight-metric">{insight.value}</div>')
                html.append(f'<div class="insight-desc">{insight.description}</div>')
                html.append('</div>')
            html.append('</div>')
            html.append('</div>')

        # Detailed Analysis
        if report.detailed_analysis:
            html.append('<div class="report-section">')
            html.append('<h3>Detailed Analysis</h3>')
            # Render internal markdown bold/italics just for the detailed analysis
            formatted_analysis = report.detailed_analysis.replace("\\n", "<br>")
            import re
            formatted_analysis = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", formatted_analysis)
            formatted_analysis = re.sub(r"\*(.+?)\*", r"<em>\1</em>", formatted_analysis)
            html.append(f'<p>{formatted_analysis}</p>')
            html.append('</div>')

        # Conclusion
        html.append('<div class="report-section">')
        html.append('<h3>Conclusion</h3>')
        html.append(f'<p>{report.conclusion}</p>')
        html.append('</div>')

        # Recommendations
        if report.recommendations:
            html.append('<div class="report-section">')
            html.append('<h3>Recommendations</h3>')
            html.append('<ol>')
            for rec in report.recommendations:
                html.append(f'<li>{rec}</li>')
            html.append('</ol>')
            html.append('</div>')

        # Data Quality
        if report.data_quality:
            html.append('<div class="report-section">')
            html.append('<h3>Data Quality & Limitations</h3>')
            html.append('<ul>')
            for dq in report.data_quality:
                html.append(f'<li>{dq}</li>')
            html.append('</ul>')
            html.append('</div>')

        html.append('</div>') # end report-document
        return "".join(html)

    @staticmethod
    def render_docx(report: StructuredReport, output_path: Path) -> Path:
        """Render the structured report to a DOCX document."""
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = docx.Document()

        # Document Styles
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        # Title
        title_para = doc.add_heading(report.title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        if report.subtitle:
            sub_para = doc.add_paragraph(report.subtitle)
            sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph() # Spacing

        # Executive Summary
        doc.add_heading('Executive Summary', level=1)
        doc.add_paragraph(report.executive_summary)

        # Objective
        doc.add_heading('Analysis Objective', level=1)
        doc.add_paragraph(report.objective)

        # Data Overview
        if report.data_overview:
            doc.add_heading('Data Overview', level=1)
            doc.add_paragraph(report.data_overview)

        # KPIs
        if report.kpis:
            doc.add_heading('Key Metrics', level=1)
            # Add a table for KPIs for better layout in Word
            table = doc.add_table(rows=1, cols=len(report.kpis))
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for i, kpi in enumerate(report.kpis):
                hdr_cells[i].text = f"{kpi.label}\n{kpi.value}"
                if kpi.description:
                    hdr_cells[i].text += f"\n{kpi.description}"
            doc.add_paragraph()

        # Key Findings
        if report.key_findings:
            doc.add_heading('Key Findings', level=1)
            for finding in report.key_findings:
                doc.add_paragraph(finding, style='List Bullet')
            doc.add_paragraph()

        # Visual Analysis
        if report.visualizations:
            doc.add_heading('Visual Analysis', level=1)
            for viz in report.visualizations:
                doc.add_heading(viz.title, level=2)
                try:
                    if viz.image_path:
                        doc.add_picture(viz.image_path, width=Inches(5.5))
                except Exception as e:
                    logger.warning("Could not insert inline chart %s: %s", viz.image_path, e)
                
                if viz.insight:
                    p = doc.add_paragraph()
                    p.add_run("Insight: ").italic = True
                    p.add_run(viz.insight)
                doc.add_paragraph()

        # Insights
        if report.insights:
            doc.add_heading('AI Insights', level=1)
            for insight in report.insights:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{insight.title}: ").bold = True
                p.add_run(insight.description)
            doc.add_paragraph()

        # Detailed Analysis
        if report.detailed_analysis:
            doc.add_heading('Detailed Analysis', level=1)
            doc.add_paragraph(report.detailed_analysis.replace("**", "").replace("*", ""))

        # Conclusion
        doc.add_heading('Conclusion', level=1)
        doc.add_paragraph(report.conclusion)

        # Recommendations
        if report.recommendations:
            doc.add_heading('Recommendations', level=1)
            for rec in report.recommendations:
                doc.add_paragraph(rec, style='List Number')
            doc.add_paragraph()

        # Data Quality
        if report.data_quality:
            doc.add_heading('Data Quality & Limitations', level=1)
            for dq in report.data_quality:
                doc.add_paragraph(dq, style='List Bullet')

        doc.save(str(output_path))
        return output_path
